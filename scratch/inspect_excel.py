import openpyxl

wb = openpyxl.load_workbook('prompts/RMO-HRA-QF-001 RMO Staff Requisition Form.xlsx')
ws = wb['Staff Requisition Form']

print('=== COLUMN WIDTHS ===')
for c in range(1, 11):
    letter = openpyxl.utils.get_column_letter(c)
    cd = ws.column_dimensions.get(letter)
    print(f'Col {letter} (#{c}): width={cd.width if cd else "None"}')

print('\n=== MERGED CELLS ===')
for m in sorted([str(m) for m in ws.merged_cells.ranges]):
    print(m)

print('\n=== ROWS ===')
for r in range(1, 75):
    h = ws.row_dimensions[r].height if r in ws.row_dimensions else None
    vals = [(openpyxl.utils.get_column_letter(c), ws.cell(r, c).value) for c in range(1, 11) if ws.cell(r, c).value is not None]
    if vals or (h is not None and h > 0):
        print(f'Row {r:2d} (h={h}): {vals}')
