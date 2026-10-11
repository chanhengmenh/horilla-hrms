"""Unit tests for the shared data import / export helpers.

These cover the CSV + Excel holiday importers, the generic Excel error-report
generator and the shared ``export_data`` Excel writer.
"""

import io

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from base.models import Holidays
from base.views import csv_holiday_import, excel_holiday_import, generate_error_report
from horilla.testkit import make_company, make_employee, make_user


def _csv_file(content: bytes) -> SimpleUploadedFile:
    return SimpleUploadedFile("holidays.csv", content, content_type="text/csv")


class CsvHolidayImportTests(TestCase):
    def test_valid_rows_after_an_invalid_row_are_still_imported(self):
        """A single bad row must not stop later valid rows being saved."""
        content = (
            b"Holiday Name,Start Date,End Date,Recurring\n"
            b"Good Before,2024-01-01,2024-01-01,yes\n"
            b"Bad Row,not-a-date,2024-01-02,maybe\n"
            b"Good After,2024-02-01,2024-02-01,no\n"
        )
        errors, total = csv_holiday_import(_csv_file(content))

        self.assertEqual(total, 3)
        # Only the bad row is an error, reported once even though it failed
        # two validations (bad date AND bad recurring flag).
        self.assertEqual(len(errors), 1)
        self.assertEqual(errors[0]["Holiday Name"], "Bad Row")
        self.assertEqual(Holidays.objects.filter(name="Good Before").count(), 1)
        self.assertEqual(
            Holidays.objects.filter(name="Good After").count(),
            1,
            "valid row after an invalid row was silently dropped",
        )

    def test_header_only_file_does_not_raise(self):
        content = b"Holiday Name,Start Date,End Date,Recurring\n"
        errors, total = csv_holiday_import(_csv_file(content))

        self.assertEqual(total, 0)
        self.assertEqual(errors, [])
        self.assertEqual(Holidays.objects.count(), 0)

    def test_malformed_row_does_not_crash_the_import(self):
        """A row with the wrong number of columns is reported, not fatal."""
        content = (
            b"Holiday Name,Start Date,End Date,Recurring\n"
            b"OnlyThree,2024-01-01,2024-01-01\n"
            b"Good,2024-03-01,2024-03-01,yes\n"
        )
        errors, total = csv_holiday_import(_csv_file(content))

        self.assertEqual(total, 2)
        self.assertEqual(len(errors), 1)
        self.assertIn("Other Errors", errors[0])
        self.assertEqual(Holidays.objects.filter(name="Good").count(), 1)


class ExcelHolidayImportTests(TestCase):
    def test_generic_exception_is_reported_under_other_errors(self):
        """Missing required column surfaces under the ``Other Errors`` key.

        The error-report template uses ``Other Errors`` (capital E), so the
        importer must use the same key or the message is silently dropped.
        """
        import pandas as pd

        df = pd.DataFrame(
            {
                "Start Date": ["2024-01-01"],
                "End Date": ["2024-01-01"],
                "Recurring": ["yes"],
            }
        )
        buf = io.BytesIO()
        df.to_excel(buf, index=False)
        buf.seek(0)

        errors, total = excel_holiday_import(buf)

        self.assertEqual(total, 1)
        self.assertEqual(len(errors), 1)
        self.assertIn(
            "Other Errors",
            errors[0],
            "excel importer used a key the error report does not recognise",
        )


class GenerateErrorReportTests(TestCase):
    def test_does_not_mutate_the_caller_template(self):
        """The shared template must survive a call unchanged.

        ``employee`` passes a module-level template; the old implementation
        deleted missing columns from it, corrupting later imports.
        """
        template = {"First": [], "Second": [], "Third": []}
        original_keys = set(template)

        path_info = generate_error_report([{"First": "x"}], template, "t.xlsx")

        self.assertTrue(path_info)
        self.assertEqual(
            set(template),
            original_keys,
            "generate_error_report deleted columns from the shared template",
        )
        for values in template.values():
            self.assertEqual(values, [])


class ExportDataTests(TestCase):
    def test_user_without_employee_record_is_redirected_to_login(self):
        """CompanyMiddleware logs out credential-only users by design.

        ``_set_company_session`` cannot resolve a company for a user with no
        linked Employee, so it logs them out and redirects to login before the
        export view ever runs. This documents that contract.
        """
        user = make_user("solo_admin", is_superuser=True)
        self.assertFalse(hasattr(user, "employee_get"))

        self.client.force_login(user)
        response = self.client.get(reverse("holiday-info-export"))

        self.assertEqual(response.status_code, 302)
        self.assertIn("login", response["Location"])

    def test_export_includes_filtered_holiday_rows(self):
        company = make_company("Export Co")
        admin = make_user("export_admin", is_superuser=True)
        make_employee(
            company=company,
            email="export_admin_emp@test.horilla",
            first_name="Export",
            last_name="Admin",
            user=admin,
        )
        Holidays.objects.create(
            name="Founders Day",
            start_date="2024-07-01",
            end_date="2024-07-01",
            recurring=False,
            company_id=company,
        )

        self.client.force_login(admin)
        response = self.client.get(
            reverse("holiday-info-export"),
            {"selected_fields": ["name", "start_date", "recurring"]},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content[:2], b"PK")

        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        self.assertIn("Name", headers)
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][headers.index("Name")], "Founders Day")
